"""Parser for Rapid Response capacity questions and Excel export."""

import json
import os
from datetime import datetime
from tempfile import NamedTemporaryFile
from typing import Any, Dict, List, Optional, Set

from django.core.cache import cache
from django.shortcuts import get_object_or_404
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from api.models import Country  
from per.ucl_research.ifrc_client import IFRCAPIClient
from per.ucl_research.blob_upload import upload_to_blob
from per.ucl_research.ops_learning_summary4 import RRCapacityTask, BaseAITask


class RapidResponseCapacityParser:
    """Coordinates data input, generated responses, and Excel export."""
    
    def __init__(self):
        self.ifrc_client = IFRCAPIClient()
        self.response_service = RRCapacityTask()
    
    def process_rr_capacity_questions_with_data(
        self, 
        country_id: int, 
        disaster_type_id: int, 
        cache_key: str,
        ops_learning_data: List[Dict[str, Any]],
        events_data: List[Dict[str, Any]]
    ) -> str:
        """Process RR capacity questions and return blob URL for the Excel file."""
        import asyncio
        
        # Load questions data
        questions_data = self._load_questions_data()
        
        # Data is provided by the caller
        
        # Process questions and fill missing fields
        processed_questions = self._process_questions(
            questions_data, events_data, ops_learning_data
        )
        
        # Generate Excel file
        filename = f"rr_capacity_filled_{country_id}_{disaster_type_id}.xlsx"
        workbook = self._create_rr_capacity_excel(
            processed_questions, country_id, disaster_type_id, 
            events_data, ops_learning_data
        )
        
        # Save and upload
        with NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            workbook.save(tmp.name)
            file_path = tmp.name
        
        blob_url = upload_to_blob(file_path, blob_name=filename)
        os.remove(file_path)
        
        # Cache result
        BaseAITask.set_cached_result(cache_key, blob_url)
        
        return blob_url
    
    @classmethod
    def _load_questions_data(cls) -> List[Dict[str, Any]]:
        """Load the parsed questions data from rr_parsed_excel.json."""
        # Get the directory of this file and join with the JSON file name
        current_dir = os.path.dirname(os.path.abspath(__file__))
        json_path = os.path.join(current_dir, 'rr_parsed_excel.json')
        with open(json_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    
    
    def _process_questions(
        self, 
        questions_data: List[Dict[str, Any]], 
        events_data: List[Dict[str, Any]], 
        ops_learning_data: List[Dict[str, Any]]
    ) -> List[Dict[str, Any]]:
        """Process each question and fill missing fields using the response service."""
        processed_questions = []
        
        # Handle None values
        if not events_data:
            events_data = []
        if not ops_learning_data:
            ops_learning_data = []
        
        for question in questions_data:
            if not isinstance(question, dict):
                continue
                
            processed_question = question.copy()
            
            # Check if the Notes on Response field is missing or null
            notes_field_new = "Notes on Response Capacity with sources"
            notes_field_old = "Notes on Response include the source"
            
            notes_value = question.get(notes_field_new) or question.get(notes_field_old)
            needs_processing = (
                not notes_value or 
                str(notes_value or "").lower() in ['nan', 'null', 'none', '']
            )
            
            if needs_processing:
                try:
                    # Use response service to fill missing fields
                    generated_responses = self.response_service.process_capacity_question(
                        question, events_data, ops_learning_data
                    )
                    
                    # Update the question with generated responses
                    for field_name, response in generated_responses.items():
                        if response:
                            if field_name == "Notes on Response Capacity with sources":
                                # Update both field names for compatibility
                                processed_question["Notes on Response Capacity with sources"] = response
                                processed_question["Notes on Response include the source"] = response
                            elif field_name in processed_question:
                                # Only update if the original field was empty/null
                                original_value = question.get(field_name)
                                if not original_value or str(original_value or "").lower() in ['nan', 'null', 'none', '']:
                                    processed_question[field_name] = response
                        
                except Exception as e:
                    # If processing fails, add error note
                    error_message = f"Processing failed: {str(e)}"
                    processed_question["Notes on Response Capacity with sources"] = error_message
                    processed_question["Notes on Response include the source"] = error_message
            
            processed_questions.append(processed_question)
        
        return processed_questions
    
    def _create_rr_capacity_excel(
        self, 
        questions_data: List[Dict[str, Any]], 
        country_id: int, 
        disaster_type_id: int, 
        events_data: Optional[List[Dict[str, Any]]] = None, 
        ops_learning_data: Optional[List[Dict[str, Any]]] = None
    ) -> Workbook:
        """Create the Excel file with filled data and sources."""
        
        wb = Workbook()
        ws = wb.active
        if ws:
            ws.title = "RR Capacity Assessment"
        else:
            ws = wb.create_sheet("RR Capacity Assessment")
        
        # Define headers
        headers = [
            "Area",
            "Critical Questions", 
            "Guiding/probing questions",
            "Notes on Response Capacity with sources",
            "Status",
            "Recommended actions for continuation of response",
            "Examples of recommended actions",
            "References"
        ]
        
        # Header styles
        header_font = Font(bold=True, size=11, color="000000")
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        
        # Main title row
        ws.merge_cells('A1:H1')
        title_cell = ws['A1']
        title_cell.value = "Rapid Response Capacity Check"
        title_cell.font = Font(bold=True, size=20, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 34
        
        # Country and date row
        country_cell = ws['A2']
        country_cell.value = "Country:"
        country_cell.font = Font(bold=True, size=20, color="000000")
        country_cell.alignment = Alignment(horizontal="left")
        
        # Get country name from country_id
        try:
            country = get_object_or_404(Country, id=country_id)
            country_name = country.name
        except:
            country_name = f"{country_id}"
        
        country_value_cell = ws['B2']
        country_value_cell.value = country_name
        country_value_cell.font = Font(size=20, color="000000")
        country_value_cell.alignment = Alignment(horizontal="left")
        
        date_cell = ws['D2']
        date_cell.value = "Date:"
        date_cell.font = Font(bold=True, size=20, color="000000")
        date_cell.alignment = Alignment(horizontal="left")
        
        # Current date
        current_date = datetime.now().strftime("%d %B %Y")
        date_value_cell = ws['E2']
        date_value_cell.value = current_date
        date_value_cell.font = Font(size=20, color="000000")
        date_value_cell.alignment = Alignment(horizontal="left")
        ws.row_dimensions[2].height = 26
        
        # Headers row
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 31
        
        # Add data rows with area-block merging
        row_idx = 4
        area_start = row_idx
        current_area = None
        merged_areas = []
        
        # Process each question
        for question in questions_data:
            if not isinstance(question, dict):
                continue
                
            area = question.get("Area", "")
            
            # If area is null/empty, use the previous area
            if not area or area == "null" or str(area or "").lower() == 'nan':
                area = current_area
            
            # Extract main area name for comparison
            if area:
                current_main_area = area.split('\n')[0].strip()
            else:
                current_main_area = None
                
            if current_area:
                previous_main_area = current_area.split('\n')[0].strip()
            else:
                previous_main_area = None
            
            # Check if we're starting a new area
            if current_main_area != previous_main_area:
                if current_area is not None:  # close previous block
                    ws.merge_cells(start_row=area_start, start_column=1,
                                    end_row=row_idx-1, end_column=1)
                    merged_areas.append({
                        'start_row': area_start,
                        'end_row': row_idx - 1,
                        'area': current_area
                    })
                area_start = row_idx
                if current_area is None or previous_main_area != current_main_area:
                    current_area = area
            
            # Write the cells
            for col_idx, header in enumerate(headers, 1):
                value = question.get(header, "")
                
                # Handle NaN values
                if str(value or "").lower() == 'nan':
                    value = ""
                
                # Handle list values
                if isinstance(value, list):
                    if value:
                        # Special handling for References column
                        if col_idx == 8:
                            reference_strings = []
                            for item in value:
                                if isinstance(item, dict):
                                    text = item.get('text', '')
                                    url = item.get('url', '')
                                    if url:
                                        reference_strings.append(f"{text} ({url})")
                                    else:
                                        reference_strings.append(text)
                                else:
                                    reference_strings.append(str(item))
                            value = '\n'.join(reference_strings)
                        else:
                            value = '\n'.join([str(item) for item in value])
                    else:
                        value = ""
                
                # Handle line breaks in text
                if value and isinstance(value, str):
                    value = value.replace('\\n\\n', '\n')
                    value = value.replace('\\n', '\n')
                
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(size=11, color="000000")
                
                # Remove markdown emphasis
                if value and isinstance(value, str) and "**" in value:
                    import re
                    
                    value = re.sub(r'\*\*-\s*([^:]+):\*\*', r'- \1:', value)
                    value = re.sub(r'-\s*\*\*([^:]+):\*\*', r'- \1:', value)
                    value = re.sub(r'\*\*([^:]+):\*\*', r'\1:', value)
                    value = re.sub(r'\*\*([^*]+)\*\*', r'\1', value)
                    value = re.sub(r'\*([^*]+)\*', r'\1', value)
                    value = re.sub(r'_([^_]+)_', r'\1', value)
                    value = value.replace('**', '')
                    
                    cell.value = value
                
                # Set row height for better readability
                ws.row_dimensions[row_idx].height = 60
            
            row_idx += 1
        
        # Merge the final block
        if current_area is not None:
            ws.merge_cells(start_row=area_start, start_column=1,
                            end_row=row_idx-1, end_column=1)
            merged_areas.append({
                'start_row': area_start,
                'end_row': row_idx - 1,
                'area': current_area
            })
            
            # Style the merged area cells
            for merged_area in merged_areas:
                area_color = self._get_area_color(merged_area['area'])
                area_cell = ws.cell(merged_area['start_row'], 1)
                area_cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                area_cell.fill = PatternFill(start_color=area_color, end_color=area_color, fill_type="solid")
                
                area_text = merged_area['area']
                if area_text and '\n' in area_text:
                    parts = area_text.split('\n', 1)
                    main_area = parts[0].strip()
                    description = parts[1].strip() if len(parts) > 1 else ""
                    full_text = f"{main_area}\n{description}".replace('\n\n', '\n')
                    area_cell.value = full_text
                    area_cell.font = Font(bold=True, size=11)
                else:
                    area_cell.value = area_text
                    area_cell.font = Font(bold=True, size=11)
                
                # Ensure all cells in the merged area are properly aligned
                for row in range(merged_area['start_row'], merged_area['end_row'] + 1):
                    cell = ws.cell(row=row, column=1)
                    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        
        # Adjust column widths
        column_widths = [25, 35, 30, 40, 15, 40, 30, 20]
        column_letters = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']
        for col_idx, width in enumerate(column_widths):
            if col_idx < len(column_letters):
                ws.column_dimensions[column_letters[col_idx]].width = width
        
        # Add Sources sheet
        self._add_sources_sheet(wb, events_data, ops_learning_data)
        
        return wb
    
    def _get_area_color(self, area: str) -> str:
        """Get a pastel color for each area."""
        if not area:
            return "FFFFFF"
        
        area_name = area.split('\n')[0].strip()
        
        area_colors = {
            "Policy, Strategy and Standards": "E6E6FA",  # Light lavender/purple
            "Analysis and Planning": "FFFACD",           # Light yellow/cream
            "Operational Capacity": "E6F3FF",           # Light blue
            "Coordination": "E6FFE6",                   # Light green
            "Operations Support": "FFE6F0"              # Light pink/rose
        }
        
        return area_colors.get(area_name, "FFFFFF")
    
    def _add_sources_sheet(
        self, 
        workbook: Workbook, 
        events_data: Optional[List[Dict[str, Any]]] = None, 
        ops_learning_data: Optional[List[Dict[str, Any]]] = None
    ):
        """Add a Sources sheet with source information."""
        sources_ws = workbook.create_sheet("Sources Used")
        
        # Styles
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        section_font = Font(bold=True, size=11, color="203764")
        
        row = 1
        
        # Title
        sources_ws.merge_cells(f'A{row}:E{row}')
        title_cell = sources_ws[f'A{row}']
        title_cell.value = "Data Sources Used in AI Analysis"
        title_cell.font = Font(bold=True, size=14, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center")
        row += 2
        
        # Events section
        if events_data:
            sources_ws[f'A{row}'] = "EVENTS"
            sources_ws[f'A{row}'].font = section_font
            row += 1
            
            event_headers = ["Event Name", "Appeal Code", "Disaster Type", "Location", "Source Context"]
            for col, header in enumerate(event_headers, 1):
                cell = sources_ws.cell(row=row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            row += 1
            
            for event in events_data[:15]:
                event_name = event.get('name', 'Unknown')
                appeal_code = event.get('appeal_source', 'N/A')
                
                dtype_info = event.get('dtype', {})
                disaster_type = dtype_info.get('name') if isinstance(dtype_info, dict) else event.get('dtype_name', 'Unknown')
                
                countries = event.get('countries', [])
                if countries:
                    first_country = countries[0]
                    if isinstance(first_country, dict):
                        location = first_country.get('name', 'Unknown')
                    else:
                        location = f"Country ID: {first_country}"
                else:
                    location = event.get('country_name', 'Unknown')
                
                source_context = event.get('source_note', 'Event from ops learning')
                
                sources_ws[f'A{row}'] = event_name
                sources_ws[f'B{row}'] = appeal_code
                sources_ws[f'C{row}'] = disaster_type
                sources_ws[f'D{row}'] = location
                sources_ws[f'E{row}'] = source_context
                row += 1
            
            row += 1
        
        # Ops Learning section
        if ops_learning_data:
            sources_ws[f'A{row}'] = "OPERATIONAL LEARNING"
            sources_ws[f'A{row}'].font = section_font
            row += 1
            
            learning_headers = ["Learning Summary", "Appeal Code", "Event", "Document", "Source Context"]
            for col, header in enumerate(learning_headers, 1):
                cell = sources_ws.cell(row=row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            row += 1
            
            for learning in ops_learning_data[:20]:
                learning_text = (
                    learning.get('learning_validated_en') or 
                    learning.get('learning_validated') or 
                    learning.get('learning_en') or 
                    'Learning content'
                )
                # Truncate if too long
                if len(learning_text) > 80:
                    learning_text = learning_text[:80] + "..."
                
                appeal_info = learning.get('appeal', {})
                if isinstance(appeal_info, dict):
                    appeal_code = appeal_info.get('code', 'N/A')
                    event_details = appeal_info.get('event_details', {})
                    if isinstance(event_details, dict):
                        event_name = event_details.get('name', 'N/A')
                    else:
                        event_name = 'N/A'
                else:
                    appeal_code = str(appeal_info) if appeal_info else 'N/A'
                    event_name = 'N/A'
                
                document = learning.get('document_name', 'N/A')
                source_context = learning.get('source_note', 'Operational learning')
                
                sources_ws[f'A{row}'] = learning_text
                sources_ws[f'B{row}'] = appeal_code
                sources_ws[f'C{row}'] = event_name
                sources_ws[f'D{row}'] = document
                sources_ws[f'E{row}'] = source_context
                row += 1
        
        # Column widths
        for col in range(1, 6):
            sources_ws.column_dimensions[chr(64 + col)].width = 25